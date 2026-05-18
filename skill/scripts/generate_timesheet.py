#!/usr/bin/env python3
"""Generate a formatted XLSX timesheet from the structured JSON entries."""

import json
import sys
import os
from datetime import datetime

def main():
    if len(sys.argv) < 3:
        print("Usage: generate_timesheet.py <entries.json> <output.xlsx>")
        sys.exit(1)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        os.system("pip install openpyxl --break-system-packages -q")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    entries_path = sys.argv[1]
    output_path = sys.argv[2]

    with open(entries_path) as f:
        data = json.load(f)

    entries = data["entries"]
    summary = data.get("summary", {})

    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"

    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='2F5496')
    date_font = Font(name='Arial', bold=True, size=10, color='2F5496')
    date_fill = PatternFill('solid', fgColor='D6E4F0')
    daily_fill = PatternFill('solid', fgColor='FFF2CC')
    lunch_fill = PatternFill('solid', fgColor='E2EFDA')
    task_font = Font(name='Arial', size=10)
    time_font = Font(name='Arial', size=10, color='555555')
    thin_border = Border(
        left=Side(style='thin', color='B4C6E7'),
        right=Side(style='thin', color='B4C6E7'),
        top=Side(style='thin', color='B4C6E7'),
        bottom=Side(style='thin', color='B4C6E7')
    )

    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 90
    ws.column_dimensions['E'].width = 10

    headers = ['Date', 'Start', 'End', 'Task', 'Hours']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[1].height = 24

    # Group entries by date
    by_date = {}
    for e in entries:
        d = e["date"]
        if d not in by_date:
            by_date[d] = []
        by_date[d].append(e)

    row = 2
    for date_str in sorted(by_date.keys()):
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        day_label = dt.strftime("%a, %b %d")

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=day_label)
        cell.font = date_font
        cell.fill = date_fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = thin_border
        for c in range(2, 6):
            ws.cell(row=row, column=c).fill = date_fill
            ws.cell(row=row, column=c).border = thin_border
        ws.row_dimensions[row].height = 22
        row += 1

        for entry in by_date[date_str]:
            ws.cell(row=row, column=1, value="").border = thin_border

            c_start = ws.cell(row=row, column=2, value=entry["start"])
            c_start.font = time_font
            c_start.alignment = Alignment(horizontal='center')
            c_start.border = thin_border

            c_end = ws.cell(row=row, column=3, value=entry["end"])
            c_end.font = time_font
            c_end.alignment = Alignment(horizontal='center')
            c_end.border = thin_border

            c_task = ws.cell(row=row, column=4, value=entry["description"])
            c_task.font = task_font
            c_task.alignment = Alignment(wrap_text=True, vertical='center')
            c_task.border = thin_border

            hours = entry.get("hours", 0)
            c_hours = ws.cell(row=row, column=5, value=hours if hours > 0 else "")
            c_hours.font = task_font
            c_hours.alignment = Alignment(horizontal='center')
            c_hours.border = thin_border

            desc_lower = entry["description"].lower()
            if "lunch" in desc_lower:
                for c in range(1, 6):
                    ws.cell(row=row, column=c).fill = lunch_fill
            elif "daily" in desc_lower and "standup" in desc_lower:
                for c in range(1, 6):
                    ws.cell(row=row, column=c).fill = daily_fill

            ws.row_dimensions[row].height = 20
            row += 1

    row += 1
    total_hours = sum(e.get("hours", 0) for e in entries if e.get("hours", 0) > 0 and "lunch" not in e.get("description", "").lower())
    ws.cell(row=row, column=4, value="Total Productive Hours").font = Font(name='Arial', bold=True, size=11)
    ws.cell(row=row, column=5, value=total_hours).font = Font(name='Arial', bold=True, size=11, color='2F5496')
    ws.cell(row=row, column=5).alignment = Alignment(horizontal='center')

    ws.freeze_panes = 'A2'
    wb.save(output_path)
    print(json.dumps({"status": "ok", "output": output_path, "total_hours": total_hours, "total_days": len(by_date)}))

if __name__ == "__main__":
    main()
